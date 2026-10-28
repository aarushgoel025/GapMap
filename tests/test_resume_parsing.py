import io
from pathlib import Path
from unittest.mock import patch
import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from resume_parsing.privacy import scrub_pii
from resume_parsing.sections import split_resume_sections
from resume_parsing.stuffing import detect_skill_stuffing, build_skill_pattern
from resume_parsing.ai_extractor import extract_resume_with_ai
from shared.schemas import ExtractionResult, Skill

client = TestClient(app)


# ---------------------------------------------------------------------------
# 1. Privacy PII Scrubbing Tests
# ---------------------------------------------------------------------------
def test_scrub_pii_removes_emails_and_phones_and_addresses():
    raw_text = """
    Candidate: Jane Doe
    Email: jane.doe@example.com
    Phone: +1 (555) 123-4567 or +91 98765 43210
    Address: 123 Main Street, Suite 4B
    
    Education: B.S. Computer Science
    Experience: Software Engineer at Acme Corp.
    """
    cleaned = scrub_pii(raw_text)

    assert "jane.doe@example.com" not in cleaned
    assert "[REDACTED_EMAIL]" in cleaned
    assert "+1 (555) 123-4567" not in cleaned
    assert "[REDACTED_PHONE]" in cleaned
    assert "123 Main Street" not in cleaned
    assert "[REDACTED_ADDRESS]" in cleaned
    # Non-PII content remains intact
    assert "B.S. Computer Science" in cleaned
    assert "Acme Corp" in cleaned


# ---------------------------------------------------------------------------
# 2. Section Splitting Tests
# ---------------------------------------------------------------------------
def test_split_resume_sections_detects_standard_headings():
    resume_text = """
    Jane Doe - Engineering Lead

    ## Technical Skills:
    Python, TypeScript, PostgreSQL, Docker

    ## Notable Projects:
    Built an AI talent assessment platform with FastAPI and Vite.

    ## Work Experience:
    Senior Software Engineer at Horizon Labs for 3 years.

    ## Education:
    B.Tech in Computer Science, 2021
    """
    sections = split_resume_sections(resume_text)

    assert "skills" in sections
    assert "Python, TypeScript" in sections["skills"]
    assert "projects" in sections
    assert "FastAPI and Vite" in sections["projects"]
    assert "experience" in sections
    assert "Horizon Labs" in sections["experience"]
    assert "education" in sections
    assert "B.Tech in Computer Science" in sections["education"]


# ---------------------------------------------------------------------------
# 3. Skill-Stuffing Detection & Boundary Tests
# ---------------------------------------------------------------------------
def test_skill_stuffing_detection_and_boundaries():
    # Boundary test: "Go" should not match "Django", "Google"
    pattern_go = build_skill_pattern("Go")
    assert pattern_go.search("Proficient in Go programming") is not None
    assert pattern_go.search("Built apps with Django and Google Cloud") is None

    # Boundary test: "C++" symbol handling
    pattern_cpp = build_skill_pattern("C++")
    assert pattern_cpp.search("Expert in C++ systems") is not None

    # Stuffing test: skill with no project/experience mentions gets flagged and confidence downgraded
    skills = ["Python", "Docker", "Kubernetes"]
    sections = {
        "skills": "Python, Docker, Kubernetes",
        "projects": "Built a scalable microservice with Python.",
        "experience": "Deployed backend containers using Docker.",
    }

    checks = detect_skill_stuffing(skills, sections, initial_confidence="HIGH")
    results = {c.name: c for c in checks}

    # Python has project evidence -> not stuffed, HIGH confidence
    assert results["Python"].stuffing_flag is False
    assert results["Python"].confidence == "HIGH"
    assert results["Python"].support_count == 1

    # Docker has experience evidence -> not stuffed, HIGH confidence
    assert results["Docker"].stuffing_flag is False
    assert results["Docker"].confidence == "HIGH"
    assert results["Docker"].support_count == 1

    # Kubernetes has zero evidence -> STUFFED, confidence downgraded to MEDIUM
    assert results["Kubernetes"].stuffing_flag is True
    assert results["Kubernetes"].confidence == "MEDIUM"
    assert results["Kubernetes"].support_count == 0


# ---------------------------------------------------------------------------
# 4. Offline AI Extractor Mock Tests
# ---------------------------------------------------------------------------
def test_gemini_path_uses_structured_output_offline():
    fake_result = ExtractionResult(
        source_type="resume",
        source_file="sample_resume.docx",
        skills=[
            Skill(
                skill_name="Python",
                category_code="COD",
                evidence="Built a Python application",
                confidence="high",
                level_estimate=7,
                source_section="Projects",
                stuffing_flag=False,
            )
        ],
        extras={
            "education": "B.Tech Computer Science",
            "projects": [],
            "experience": [],
            "certifications": [],
        },
    )

    with (
        patch("resume_parsing.ai_extractor.get_gemini") as mock_gemini,
        patch("resume_parsing.ai_extractor.call_structured", return_value=fake_result) as mock_call,
    ):
        mock_gemini.return_value = object()

        result = extract_resume_with_ai(
            resume_text="Built a Python application",
            source_file="sample_resume.docx",
            provider="gemini",
        )

    assert result.source_type == "resume"
    assert result.source_file == "sample_resume.docx"
    assert result.skills[0].skill_name == "Python"
    mock_gemini.assert_called_once()
    mock_call.assert_called_once()


def test_mistral_path_uses_structured_output_offline():
    fake_result = ExtractionResult(
        source_type="resume",
        source_file="sample_resume.docx",
        skills=[],
        extras={},
    )

    with (
        patch("resume_parsing.ai_extractor.get_mistral") as mock_mistral,
        patch("resume_parsing.ai_extractor.call_structured", return_value=fake_result) as mock_call,
    ):
        mock_mistral.return_value = object()

        result = extract_resume_with_ai(
            resume_text="Python developer",
            source_file="sample_resume.docx",
            provider="mistral",
        )

    assert result.source_type == "resume"
    mock_mistral.assert_called_once()
    mock_call.assert_called_once()


# ---------------------------------------------------------------------------
# 5. FastAPI Resume Endpoints Tests
# ---------------------------------------------------------------------------
def test_resume_endpoints_status_and_list():
    res_status = client.get("/api/resume/status")
    assert res_status.status_code == 200
    assert res_status.json() == {"module": "resume_parsing", "status": "initialized"}

    res_list = client.get("/api/resume/resumes")
    assert res_list.status_code == 200
    resumes = res_list.json()
    assert len(resumes) >= 3
    assert resumes[0]["candidate_name"] == "Aarush Sharma"


def test_resume_parse_json_sample():
    res = client.post(
        "/api/resume/parse",
        json={"resume_source": "aarush_sharma_cv.pdf"},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["source_type"] == "resume"
    assert len(data["skills"]) > 0
    # Check that stuffing flag is present
    has_stuffing_field = any("stuffing_flag" in s for s in data["skills"])
    assert has_stuffing_field is True


def test_resume_parse_upload_docx():
    # Create an in-memory docx using python-docx
    import docx
    doc = docx.Document()
    doc.add_heading("Alex Morgan - Resume", level=1)
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, PostgreSQL, Docker, Kubernetes")
    doc.add_heading("Projects", level=2)
    doc.add_paragraph("Built an API in Python using PostgreSQL for database storage.")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Software Engineer at CloudTech. Deployed microservices with Docker.")
    
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    res = client.post(
        "/api/resume/parse",
        files={"file": ("alex_morgan_cv.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["source_type"] == "resume"
    assert data["source_file"] == "alex_morgan_cv.docx"
    
    skill_names = [s["skill_name"].lower() for s in data["skills"]]
    assert "python" in skill_names
    assert "postgresql" in skill_names
    assert "docker" in skill_names

    # Check stuffing detection: Kubernetes had 0 mentions in projects and experience
    k8s = next((s for s in data["skills"] if s["skill_name"].lower() == "kubernetes"), None)
    if k8s:
        assert k8s["stuffing_flag"] is True


def test_verify_github_endpoint():
    res = client.post(
        "/api/resume/verify-github",
        json={
            "github_url": "https://github.com/aarushgoel025",
            "skills": [
                {
                    "skill_name": "Python",
                    "category_code": "COD",
                    "confidence": "high",
                }
            ],
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert "verified_skills" in data
    assert "github" in data["verified_skills"][0]["verified_by"]
