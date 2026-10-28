"""Small helpers for the JD Analytics tests: tiny PDF/DOCX builders and fake model answers."""
from pathlib import Path

import docx

from app.jd_analytics.pipeline import LLMExtraction, LLMSkill

JD_LINES = [
    "Software Engineer, Backend",
    "Minimum qualifications: 3 years of Python and SQL experience.",
    "Preferred: Kubernetes and Docker experience.",
    "You will design scalable distributed systems and work with the team.",
]


def make_pdf(path: Path, lines: list[str]) -> Path:
    """Write a minimal one-page text PDF (no extra libraries needed)."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    path.write_bytes(out)
    return path


def make_docx(path: Path, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> Path:
    """Write a .docx with paragraphs, then an optional table."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    document.save(str(path))
    return path


def llm_skill(name: str, category: str = "COD", importance: str = "must_have",
              confidence: str = "high", evidence: str = "quote") -> LLMSkill:
    return LLMSkill(skill_name=name, category_code=category, importance=importance,
                    confidence=confidence, evidence=evidence)


def extraction(skills: list[LLMSkill], company: str = "Acme", role: str = "Engineer") -> LLMExtraction:
    return LLMExtraction(company=company, role=role, skills=skills)
